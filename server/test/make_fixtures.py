"""Generates real sample files for the end-to-end tests (run once: python3 make_fixtures.py)."""
import os, random, io
from reportlab.lib.pagesizes import A4, letter, landscape, A5
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from PIL import Image
import docx, openpyxl, pptx
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter

out = os.path.join(os.path.dirname(__file__), 'fixtures')
os.makedirs(out, exist_ok=True)
P = lambda n: os.path.join(out, n)
random.seed(7)

def text_pdf(path, tag, pages, size):
    c = canvas.Canvas(path, pagesize=size)
    for i in range(1, pages + 1):
        c.setFont('Helvetica-Bold', 28); c.drawString(60, size[1] - 100, f'{tag}-PAGE-{i}')
        c.setFont('Helvetica', 12)
        for l in range(30): c.drawString(60, size[1] - 150 - l * 16, f'{tag} page {i} line {l} lorem ipsum dolor sit amet')
        c.showPage()
    c.save()

text_pdf(P('a.pdf'), 'ALPHA', 3, A4)
text_pdf(P('b.pdf'), 'BRAVO', 2, landscape(letter))

# Marker PDF: black square top-left, to verify rotation/orientation geometry by rendering.
c = canvas.Canvas(P('marker.pdf'), pagesize=A5)
c.setFillColorRGB(0, 0, 0); c.rect(20, A5[1] - 80, 60, 60, fill=1)
c.setFont('Helvetica', 14); c.drawString(100, A5[1] - 60, 'MARKER'); c.showPage(); c.save()

# Heavy PDF: 4 pages of high-resolution noisy photos (hard to compress, realistic scan-like payload).
def noisy(w, h, seed):
    r = random.Random(seed)
    base = Image.effect_noise((w, h), 60).convert('RGB')
    grad = Image.linear_gradient('L').resize((w, h)).convert('RGB')
    return Image.blend(base, grad, 0.4)
c = canvas.Canvas(P('heavy.pdf'), pagesize=A4)
for i in range(4):
    img = noisy(2480, 3508, i); buf = io.BytesIO(); img.save(buf, 'JPEG', quality=92); buf.seek(0)
    c.drawImage(ImageReader(buf), 0, 0, width=A4[0], height=A4[1]); c.setFont('Helvetica-Bold', 40)
    c.setFillColorRGB(1, 1, 1); c.drawString(60, 760, f'HEAVY-{i+1}'); c.showPage()
c.save()

# Vector-heavy PDF (no images): a dense drawing like a map/CAD plan. Image compression cannot shrink it.
random.seed(3)
c = canvas.Canvas(P('vector.pdf'), pagesize=A4)
for pg in range(6):
    c.setFont('Helvetica-Bold', 24); c.drawString(50, 800, f'VECTOR-PAGE-{pg+1}')
    for _ in range(2500):
        x, y = random.uniform(30, 560), random.uniform(30, 760)
        path = c.beginPath(); path.moveTo(x, y)
        for _ in range(12): x += random.uniform(-25, 25); y += random.uniform(-25, 25); path.lineTo(x, y)
        c.setStrokeColorRGB(random.random(), random.random(), random.random()); c.drawPath(path, stroke=1, fill=0)
    c.showPage()
c.save()

d = docx.Document(); d.add_heading('WORD-DOC-TITLE', 1)
for i in range(60): d.add_paragraph(f'Word paragraph {i}: the quick brown fox jumps over the lazy dog.')
d.save(P('c.docx'))

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Sheet1'
ws.column_dimensions['A'].width = 24
ws.append(['EXCEL-HEADER', 'Value', 'Squared'])
for i in range(1, 40): ws.append([f'row{i}', i, i * i])
wb.save(P('d.xlsx'))

pr = pptx.Presentation(); lay = pr.slide_layouts[1]
for i in range(1, 4):
    s = pr.slides.add_slide(lay); s.shapes.title.text = f'PPT-SLIDE-{i}'; s.placeholders[1].text = f'Body of slide {i}'
pr.save(P('e.pptx'))

Image.effect_noise((1600, 1200), 50).convert('RGB').save(P('img1.jpg'), quality=90)
Image.new('RGB', (900, 1400), (30, 90, 200)).save(P('img2.png'))

open(P('corrupt.pdf'), 'wb').write(b'%PDF-1.7\n1 0 obj\n<< /Type /Catalog >>\nendobj\nGARBAGE' + os.urandom(500))
open(P('fake.pdf'), 'wb').write(b'MZ\x90\x00\x03\x00\x00\x00 this is not a pdf at all' * 20)
open(P('notes.txt'), 'w').write('plain text')
open(P('corrupt.docx'), 'wb').write(b'PK\x03\x04' + os.urandom(2000))

w = PdfWriter(); r = PdfReader(P('a.pdf'))
for pg in r.pages: w.add_page(pg)
w.encrypt('secret'); w.write(open(P('encrypted.pdf'), 'wb'))

for n in sorted(os.listdir(out)): print(f'{n:16} {os.path.getsize(P(n)):>10,} bytes')
